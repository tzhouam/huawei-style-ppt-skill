#!/usr/bin/env python3
"""Structural preflight for an editable PPTX; render and inspect slides as well."""

import argparse
import re
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt


PAGE_RE = re.compile(r"^\s*(\d{1,3})\s*/\s*(\d{1,3})\s*$")


def inspect(path: Path) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []
    with ZipFile(path) as archive:
        bad = archive.testzip()
        if bad:
            return [f"corrupt archive member: {bad}"], []
        media = [name for name in archive.namelist() if name.startswith("ppt/media/")]

    deck = Presentation(path)
    total = len(deck.slides)
    tolerance = Pt(2)
    for number, slide in enumerate(deck.slides, 1):
        page_marks = []
        full_images = 0
        editable = 0
        for shape in slide.shapes:
            if (shape.left < -tolerance or shape.top < -tolerance or
                    shape.left + shape.width > deck.slide_width + tolerance or
                    shape.top + shape.height > deck.slide_height + tolerance):
                warnings.append(f"slide {number}: shape bounds extend outside page: {shape.name}")
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                if (shape.width * shape.height >= deck.slide_width * deck.slide_height * .9):
                    full_images += 1
            elif shape.has_text_frame or shape.shape_type in {
                MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.LINE, MSO_SHAPE_TYPE.TABLE,
            }:
                editable += 1
            if shape.has_text_frame:
                mark = PAGE_RE.fullmatch(shape.text.strip())
                if mark:
                    page_marks.append((int(mark.group(1)), int(mark.group(2))))
        if full_images and editable <= 2:
            warnings.append(f"slide {number}: looks like a full-page image; verify editability")
        if page_marks and (number, total) not in page_marks:
            warnings.append(f"slide {number}: page marker {page_marks} differs from {number}/{total}")
        if len(page_marks) > 1:
            warnings.append(f"slide {number}: multiple page markers {page_marks}")
    if not media:
        warnings.append("no embedded media found; verify any expected screenshots/video")
    return errors, warnings


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx", type=Path)
    args = parser.parse_args()
    try:
        errors, warnings = inspect(args.pptx)
    except Exception as exc:
        print(f"ERROR: could not inspect {args.pptx}: {exc}")
        return 1
    for item in errors:
        print("ERROR:", item)
    for item in warnings:
        print("WARN:", item)
    if not errors:
        print(f"OK: {args.pptx} (structural preflight; visual QA still required)")
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
